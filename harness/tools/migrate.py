#!/usr/bin/env python3
"""Converte os Excels de pinouts antigos (uma pasta = um harness) no formato standard LART: um fio = uma linha.
O que não bate certo (pinos com duas ligações, destinos desconhecidos, nomes duplicados…) vai para a folha Issues
de cada harness. Gera também index.json, que o site usa para listar os harnesses.

Uso:  python migrate.py <pasta com as pastas dos harnesses> <pasta de saída>
      ex.: python migrate.py ../data/t26/legacy ../data/t26/std
Requer: openpyxl
"""
import json
import re
import sys
from collections import Counter, OrderedDict, defaultdict, deque
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).parent))  # config_accu ao lado, mesmo com python -I
import config_accu  # noqa: E402
import config_drawings  # noqa: E402
import config_parts  # noqa: E402

CAR = 'T26e'
HUB = 'CAR'  # harness principal: as caixas ligam-se a ele por conectores com o mesmo nome
# pasta -> (código, nome, categoria, prefixos de interface). car = no carro; offboard/test aparecem no site mas não
# entram no trace do carro. Prefixos de interface: conectores da caixa que ligam ao harness principal (HUB) com o
# mesmo nome — só esses se acoplam (TSAL J1 ou FAN J1 existem em vários sítios e são fichas diferentes).
HARNESSES = {
    'Accu_Harness': ('ACCU', 'Acumulador', 'car', ['ACCU']),
    'Car_Harness': ('CAR', 'Car harness (principal)', 'car', []),
    'LVBox_Harness': ('LVBOX', 'LV Box', 'car', ['LVBOX']),
    'JunctionBox_Harness': ('JBOX', 'Junction Box', 'car', ['JBOX']),
    'Data_Box': ('DATABOX', 'Data Box', 'car', ['DATABOX']),
    'ADBox_Harness': ('ADBOX', 'AD Box', 'car', ['ADBOX']),
    'Side Panel': ('SIDEPANEL', 'Side Panel', 'car', ['SIDEPANEL']),
    'Pedal_Box': ('PEDALBOX', 'Pedal Box', 'car', ['PEDALBOX']),
    'Aux_Front_Data': ('AUXFRONT', 'Aux Front Data', 'car', []),
    'Acquisition_2': ('ACQ2', 'Acquisition 2', 'car', ['ACQUISITION 2']),
    'Acquisition 7': ('ACQ7', 'Acquisition 7', 'car', ['ACQUISITION 7']),
    'Handcart': ('HANDCART', 'Handcart (carregamento)', 'offboard', []),
}
# pastas com mais de um Excel: um harness por ficheiro
FILES = {
    'Car_Harnes_DASH_RASPBOX_Pinouts.xlsx': ('TEST-CAR', 'Teste: car harness DASH/RASPBOX', 'test', []),
    'DataBox_DASHBOX_Harness_Pinout.xlsx': ('TEST-DASHBOX', 'Teste: DataBox DASHBOX', 'test', []),
}
CONFIGS = {'Accu_Harness': config_accu.config, 'ADBox_Harness': config_drawings.adbox,
           'Side Panel': config_drawings.sidepanel, 'Handcart': config_drawings.handcart}
# gralhas de destino que valem para todos os harnesses (a nota do fio guarda o nome original)
DEST_ALIAS = {'GOUND POINT': 'GROUND POINT'}


class P(str):
    """Valor proposto (não confirmado) — escrito a itálico laranja."""


IEC = ['BK', 'BN', 'RD', 'OG', 'YE', 'GN', 'BU', 'VT', 'GY', 'WH', 'PK', 'TQ']
COLOURS = [('BK', 'preto', ''), ('BN', 'castanho', 'GND'), ('RD', 'vermelho', 'CAN H'), ('OG', 'laranja', 'HV (obrigatório)'),
           ('YE', 'amarelo', '+24 V'), ('GN', 'verde', 'CAN carregador H'), ('BU', 'azul', 'SDC / CAN L'),
           ('VT', 'violeta', 'CAN carregador L'), ('GY', 'cinzento', ''), ('WH', 'branco', 'sinal'),
           ('PK', 'rosa', ''), ('TQ', 'turquesa', '')]
CSA_AWG = [(0.14, 26), (0.22, 24), (0.35, 22), (0.5, 20), (0.75, 18), (1, 17), (1.5, 16), (2.5, 14), (4, 12),
           (6, 10), (10, 8), (16, 6), (25, 4), (35, 2), (50, 1)]
COLOUR_ALIAS = {'blue': 'BU', 'bue': 'BU', 'azul': 'BU', 'broown': 'BN', 'brown': 'BN', 'castanho': 'BN',
                'prrpul': 'VT', 'purple': 'VT', 'roxo': 'VT', 'violet': 'VT', 'white': 'WH', 'branco': 'WH',
                'yellow': 'YE', 'amarelo': 'YE', 'orange': 'OG', 'laranja': 'OG', 'red': 'RD', 'vermelho': 'RD',
                'green': 'GN', 'verde': 'GN', 'black': 'BK', 'preto': 'BK', 'grey': 'GY', 'gray': 'GY',
                'cinzento': 'GY', 'pink': 'PK', 'rosa': 'PK'}
DIR_ALIAS = {'IN': 'IN', 'OUT': 'OUT', 'VCC': 'VCC', 'VEE': 'GND', 'NC': 'NC', 'HIGH': 'BIDIR', 'LOW': 'BIDIR'}
STD_COLS = {'Conector', 'Pin', 'Label', 'In/Out', 'Voltage', 'Wire Color', 'Wire Name', 'Wire Size', 'Destination',
            'Destination Pin', 'Description'}
UNKNOWN_DEST = {'?', '??', 'SOMEWHERE', 'NA', 'N/A', '-', 'TBD', 'X'}
BARE = re.compile(r'^[A-Z]{1,3}\d{1,2}[A-Z]?$', re.I)  # J1, D2, X1, XK1A, J10: precisa do nome do dispositivo


# ---------------------------------------------------------------------------- normalização
def s(v):
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return '' if v is None else re.sub(r'\s+', ' ', str(v)).strip()


def norm_pin(p):
    p = s(p)
    if 'interlock' in p.lower():
        return 'IL+' if '+' in p else 'IL-'
    return '' if p == '-' else p


def display(name):
    """Nome de um conector como aparece no site: maiúsculas, espaços simples, 'LVBOXJ2' -> 'LVBOX J2'."""
    n = re.sub(r'\s+', ' ', s(name).replace('_', ' ')).upper()
    return re.sub(r'(?<=[A-Z])(J\d+)$', r' \1', n)


def canon(name):
    """Chave para comparar nomes escritos de maneiras diferentes (LvBox J1 = LV_Box J1 = LVBoxJ1)."""
    n = re.sub(r'(?<=[A-Z0-9])-(?=[A-Z0-9])', '', display(name))  # DC-DC = DCDC, mas "AIR- J1" fica
    return re.sub(r'[\s.]', '', n)


def auto_device(sheet):
    d = re.sub(r'\s+J\d+$', '', sheet.strip(), flags=re.I).upper()
    return re.sub(r'^ACQUISITION BOARD (\d+)$', r'ACQUISITION \1', d)


def auto_ref(device, conn):
    return display(f'{device} {conn}' if BARE.match(conn.replace(' ', '')) and device else conn)


def colour_of(c):
    parts = [p.strip() for p in s(c).split('/') if p.strip()]
    return '/'.join(p.upper() if p.upper() in IEC else COLOUR_ALIAS.get(p.lower(), p.upper()) for p in parts)


def volt_of(v):
    v = s(v).replace('v', 'V')
    return v + 'V' if re.fullmatch(r'\d+(\.\d+)?', v) else v


def csa_of(v):
    v = s(v).lower().replace(',', '.')
    m = re.fullmatch(r'(\d+(?:\.\d+)?)\s*(mm\^?2|mm²|mm2|mm)?', v)
    if m:
        return float(m.group(1))
    m = re.fullmatch(r'(\d+)\s*awg', v)
    return next((mm for mm, awg in CSA_AWG if awg == int(m.group(1))), None) if m else None


def hv_of(colour, volts):
    return 'Y' if colour.split('/')[0] == 'OG' or any(int(n) >= 60 for v in volts for n in re.findall(r'\d+', v)) else 'N'


def net_of(label):
    n = re.sub(r'\s+', '_', label.strip().upper()).replace('_-_', '-_')  # "AIR - Inten" -> "AIR-_INTEN"
    return re.sub(r'_+', '_', n)


def ground_like(ref):
    return bool(re.search(r'GROUND|GOUND|GND|CHASSI', canon(ref)))


def levenshtein(a, b):
    d = list(range(len(b) + 1))
    for i, x in enumerate(a, 1):
        prev, d[0] = d[0], i
        for j, y in enumerate(b, 1):
            prev, d[j] = d[j], min(d[j] + 1, d[j - 1] + 1, prev + (x != y))
    return d[-1]


# ---------------------------------------------------------------------------- leitura
def read_workbook(path, cfg):
    """-> pins {(ref,pin): info}, edges [ligação escrita numa linha], notes, problems, defined {ref: {folhas}}."""
    wb = openpyxl.load_workbook(path, data_only=True)
    pins, raw, notes, problems = OrderedDict(), [], [], []
    defined = OrderedDict()
    sheets_cfg = cfg.get('sheets', {})
    for ws in wb.worksheets:
        sheet = ws.title.strip()
        if re.search(r'_old$', sheet, re.I):
            problems.append(('info', sheet, 'folha antiga (_old) — não lida'))
            continue
        if sheet.lower().startswith('nota'):
            for row in ws.iter_rows(values_only=True):
                t = ' '.join(s(v) for v in row if s(v))
                if t:
                    notes.append((sheet, t))
            continue
        groups, found = [], False
        for row in ws.iter_rows(values_only=True):
            r = [s(v) for v in row]
            if 'Conector' in r and 'Pin' in r:  # cabeçalho; pode haver duas tabelas lado a lado
                found = True
                starts = [i for i, v in enumerate(r) if v == 'Conector']
                old = {g['start']: g for g in groups}
                groups = []
                for n, i in enumerate(starts):
                    end = starts[n + 1] if n + 1 < len(starts) else len(r)
                    cols = {}
                    for j in range(i, end):
                        cols.setdefault(r[j], j) if r[j] else None
                    g = old.get(i, dict(start=i, block=-1))
                    groups.append(dict(start=i, cols=cols, block=g['block'] + 1, conn=None, last_pin=None,
                                       last=dict(colour='', volt='', desc='')))
                continue
            for g in groups:
                read_row(r, g, sheet, sheets_cfg, cfg, pins, raw, defined)
        if not found:
            problems.append(('info', sheet, 'folha sem tabela Conector/Pin — não lida'))
    for k, (label, dr, issue) in cfg.get('pin_fix', {}).items():
        pins[k].update(label=label, dir=dr, issue=issue)
    edges = resolve_dests(raw, defined, cfg, problems)
    return pins, edges, notes, problems, defined


def read_row(r, g, sheet, sheets_cfg, cfg, pins, raw, defined):
    col = lambda name: r[g['cols'][name]] if name in g['cols'] and g['cols'][name] < len(r) else ''
    conn = col('Conector') or g['conn']
    g['conn'] = conn
    pin, label, dr = norm_pin(col('Pin')), col('Label'), col('In/Out')
    volt, colour, dest, desc = volt_of(col('Voltage')), col('Wire Color'), col('Destination'), col('Description')
    size = col('Wire Size') or col('Wire Name')
    csa = csa_of(size)
    name = '' if csa is not None else col('Wire Name')
    extra = '; '.join(f'{h}: {r[j]}' for h, j in g['cols'].items() if h not in STD_COLS and j < len(r) and r[j])
    if not conn or not (pin or label or dest):
        return
    dev = sheets_cfg.get(sheet, auto_device(sheet) if 'ref_of' not in cfg else None)
    dev = dev[g['block']] if isinstance(dev, list) else dev
    if not pin and not label:  # linha de continuação: mesmo pino, outro destino
        pin = g['last_pin']
        colour, volt, desc = colour or g['last']['colour'], volt or g['last']['volt'], desc or g['last']['desc']
    pin = cfg['row_fix'](pin, label) if 'row_fix' in cfg else pin
    ref = cfg['ref_of'](dev, conn, pin) if 'ref_of' in cfg else auto_ref(dev, conn)
    if dr.upper() in ('HIGH', 'LOW') and 'can' in label.lower():
        label = f"{label}_{'H' if dr.upper() == 'HIGH' else 'L'}"
    if pin:
        defined.setdefault(ref, set()).add(sheet)
    if (ref, pin) not in pins:
        pins[(ref, pin)] = dict(label=label, dir=DIR_ALIAS.get(dr.upper(), dr.upper()), volt=volt, sheet=sheet, issue='')
    g['last_pin'], g['last'] = pin, dict(colour=colour, volt=volt, desc=desc)
    if dest:
        raw.append(dict(a=(ref, pin), dest=dest, dpin=r[g['cols']['Destination Pin']] if 'Destination Pin' in g['cols'] and g['cols']['Destination Pin'] < len(r) else '',
                        colour=colour_of(colour), volt=volt, desc=desc, sheet=sheet, csa=csa, name=name, note=extra))


def resolve_dests(raw, defined, cfg, problems):
    by_canon = OrderedDict()
    for ref in defined:
        by_canon.setdefault(canon(ref), ref)
    alias, default_pin = {**DEST_ALIAS, **cfg.get('dest_alias', {})}, cfg.get('dest_default_pin', {})
    edges = []
    for e in raw:
        up = s(e['dest']).upper()
        parts = [p for p in re.split(r'\s*\+\s*', e['dest'])]
        if up in alias:
            ref = alias[up]
            e = dict(e, note='; '.join(x for x in (e['note'], f'no Excel original: "{e["dest"]}"') if x))
        elif canon(e['dest']) in UNKNOWN_DEST:
            problems.append(('unknown', e, f"destino \"{e['dest']}\""))
            continue
        elif len(parts) > 1 and all(len(re.sub(r'\W', '', p)) >= 3 for p in parts):
            problems.append(('multi', e, f"vários destinos numa célula: \"{e['dest']}\""))
            continue
        else:
            c = canon(e['dest'])
            # "JBox" sem número: se só houver um conector desse dispositivo (JBOX J1), é esse
            only = [r for k, r in by_canon.items() if k.startswith(c) and re.fullmatch(r'J\d+', k[len(c):])]
            ref = by_canon.get(c) or (only[0] if len(only) == 1 else display(e['dest']))
        pin = norm_pin(e['dpin']) or default_pin.get(ref, '') or ('GND' if ground_like(ref) and ref not in defined else '')
        edges.append(dict(e, b=(ref, pin)))
    return edges


# ---------------------------------------------------------------------------- fios
def build_wires(pins, edges, cfg, problems, studs):
    drop = cfg.get('drop', {})
    # a cor de uma ligação descartada continua válida para o pino de onde saía (ex.: AMS J14 em I-04)
    dropped_colour = {e['a']: (e['colour'], e['sheet']) for e in edges if (e['a'], e['b']) in drop and e['colour']}
    edges = [e for e in edges if (e['a'], e['b']) not in drop]
    # destino sem pino: se a outra ponta escreveu a ligação de volta, o pino vem de lá; senão ?1, ?2… (um por fio,
    # para não juntar fios diferentes no mesmo circuito)
    back = defaultdict(list)
    for e in edges:
        back[(e['a'][0], e['b'])].append(e['a'][1])
    unknown = Counter()
    for e in edges:
        ref, pin = e['b']
        if pin:
            continue
        found = [p for p in back.get((ref, e['a']), []) if p]
        if len(found) == 1:
            e['b'] = (ref, found[0])
        else:
            unknown[ref] += 1
            e['b'] = (ref, f'?{unknown[ref]}')
            problems.append(('nopin', e, f"destino {ref} sem pino"))
    # splices: um pino com vários destinos na mesma folha
    dests = defaultdict(list)
    for e in edges:
        dests[(e['sheet'], e['a'])].append(e['b'])
    fanout = {a: set(bs) for (_, a), bs in dests.items() if len(set(bs)) > 1 and a[0] not in studs}
    splice = {a: f"SPL {cfg.get('splice_net', {}).get(a, a[0] + '-' + a[1])}" for a in fanout}
    split = []
    for e in edges:
        if e['a'] in splice:
            spl = (splice[e['a']], 'S')
            split += [dict(e, b=spl, trunk=True), dict(e, a=spl)]
        elif e['b'] in splice and e['a'] in fanout[e['b']]:
            split.append(dict(e, b=(splice[e['b']], 'S')))
        elif e['b'] in splice and cfg.get('strict'):
            raise SystemExit(f"conflito não resolvido no splice: {e['a']} -> {e['b']}")
        else:
            split.append(e)
    groups = OrderedDict()
    for e in split:
        groups.setdefault(frozenset((e['a'], e['b'])), []).append(e)
    for spl in splice.values():
        pins[(spl, 'S')] = dict(label=spl[4:], dir='', volt='', sheet='', issue='')

    order = cfg.get('priority') or list(OrderedDict.fromkeys(p['sheet'] for p in pins.values() if p['sheet']))
    rank = {sh: i for i, sh in enumerate(order)}
    net_override, issue_pins = cfg.get('net_override', {}), cfg.get('issue_pins', {})
    wires = []
    for key, es in groups.items():
        a, b = es[0]['a'], es[0]['b']
        for end in (a, b):
            pins.setdefault(end, dict(label='', dir='', volt='', sheet='', issue=''))
        # orientação: tronco = pino -> splice; ramo = splice -> pino; resto = quem fornece o sinal primeiro
        is_spl = lambda p: p[0].startswith('SPL')
        is_src = lambda p: pins[p]['dir'] in ('OUT', 'VCC', 'GND')
        trunk = any(e.get('trunk') for e in es)
        if is_spl(a) and trunk or is_spl(b) and not trunk or not (is_spl(a) or is_spl(b)) and is_src(b) and not is_src(a):
            a, b = b, a
        best = min(es, key=lambda e: rank.get(e['sheet'], 99))
        colours = list(OrderedDict.fromkeys(e['colour'] for e in es if e['colour']))
        note = f"cores diferentes nas 2 pontas: {colours}" if len(colours) > 1 else ''
        if not colours and (a in dropped_colour or b in dropped_colour):
            colour, sh = dropped_colour.get(a) or dropped_colour[b]
            colours, note = [colour], f"cor vinda da folha {sh} (ligação descartada, ver issue)"
        both = {e['a'] for e in es} == set(key) and len(key) == 2
        if is_spl(a) or is_spl(b):
            net = a[0][4:] if is_spl(a) else b[0][4:]
        else:
            lab = pins[a]['label'] or pins[b]['label']
            net = net_of(lab)
            net = net_override.get(net, net) or net_of(f"{a[0].split()[0]}_{lab}")
        volts = [pins[a]['volt'], pins[b]['volt']] + [e['volt'] for e in es]
        issues = {i for i, ps in issue_pins.items() if a in ps or b in ps} | {pins[p]['issue'] for p in (a, b)} - {''}
        csa = next((e['csa'] for e in es if e['csa'] is not None), None)
        wires.append(dict(net=net, a=a, b=b, colour=colours[0] if colours else '', colour_note=note,
                          hv=hv_of(colours[0] if colours else '', volts), volt=next((v for v in volts if v), ''),
                          desc=best['desc'], csa=csa, name=next((e['name'] for e in es if e['name']), ''),
                          note='; '.join(OrderedDict.fromkeys(e['note'] for e in es if e['note'])),
                          origin='2 pontas' if both else f"1 ponta (folha {es[0]['sheet']})",
                          sheets={e['sheet'] for e in es}, issues=issues))
    # pares entrançados propostos (CAN, USB dados)
    tp = OrderedDict()
    for w in wires:
        m = re.match(r'(.*(CAN|USB))_[HLPN]$', w['net'])
        if m:
            w['twist'] = tp.setdefault((m.group(1), tuple(sorted((w['a'][0], w['b'][0])))), f"TP{len(tp) + 1}")
    return wires


def self_check(edges, wires, cfg, studs):
    """Harness curado (strict): falha se o Excel de origem mudou e as correções deixaram de bater certo."""
    seen = {(e['a'], e['b']) for e in edges}
    stale = [k for k in cfg.get('drop', {}) if k not in seen]
    assert not stale, f"DROP sem ligação correspondente no Excel: {stale}"
    use = Counter(p for w in wires for p in (w['a'], w['b']))
    multi = [p for p, n in use.items() if n > 1 and not p[0].startswith('SPL') and p[0] not in studs]
    assert not multi, f"pinos com mais de um fio: {multi}"


def auto_issues(pins, wires, problems, defined, studs):
    """Issues geradas sem intervenção humana, para os harnesses que ainda não foram revistos."""
    out = []

    def add(sev, where, problem, a='', b='', applied='', wires_=()):
        iid = f'I-{len(out) + 1:02d}'
        out.append((iid, sev, where, problem, a, b, applied))
        for w in wires_:
            w['issues'].add(iid)

    side = lambda w, p: w['b'] if w['a'] == p else w['a']
    use = defaultdict(list)
    for w in wires:
        for p in (w['a'], w['b']):
            use[p].append(w)
    for p, ws in use.items():
        if len(ws) > 1 and not p[0].startswith('SPL') and p[0] not in studs:
            desc = [f"{':'.join(side(w, p))} (folha {', '.join(sorted(w['sheets']))})" for w in ws]
            add('crítico', ':'.join(p), f"O pino {':'.join(p)} tem {len(ws)} ligações diferentes: {'; '.join(desc)}.",
                desc[0], desc[1], 'Ficaram todas — decidir qual é a certa e apagar as outras.', ws)
    for ref, sheets in defined.items():
        if len(sheets) > 1:
            add('importante', ref, f"O conector {ref} está definido em várias folhas: {', '.join(sorted(sheets))}. Copiado por engano?")
    for kind, sev, label in (('unknown', 'importante', 'Destino desconhecido'), ('multi', 'importante', 'Vários destinos numa célula'),
                             ('nopin', 'menor', 'Destino sem pino')):
        rows = [(e, msg) for k, e, msg in problems if k == kind]
        if rows:
            hit = [w for w in wires for e, _ in rows if e['a'] in (w['a'], w['b'])] if kind == 'nopin' else []
            add(sev, f'{len(rows)} linha(s)', f"{label}: " + '; '.join(f"{':'.join(e['a'])} → {msg} (folha {e['sheet']})" for e, msg in rows[:25])
                + (' …' if len(rows) > 25 else ''), applied='Pino marcado "?"' if kind == 'nopin' else 'Linha sem fio', wires_=hit)
    refs = sorted({r for r, _ in pins if not r.startswith('SPL')})
    skeleton = lambda r: re.sub(r'[\d+\-]', '#', canon(r))  # ACQUISITION 1 J1 vs 7 J1, TS + BJ vs TS - BJ: nomes diferentes de propósito
    near = [(x, y) for i, x in enumerate(refs) for y in refs[i + 1:]
            if skeleton(x) != skeleton(y) and 0 < levenshtein(canon(x), canon(y)) <= 1]
    if near:
        add('menor', 'nomes', 'Nomes muito parecidos (gralha?): ' + '; '.join(f'{x} ~ {y}' for x, y in near[:30]))
    implicit = sorted({r for r, _ in pins if r not in defined and not r.startswith('SPL')})
    if implicit:
        add('info', 'destinos', f"{len(implicit)} conector(es) só aparecem como destino, sem folha própria: {', '.join(implicit)}.")
    one = [w for w in wires if w['origin'] != '2 pontas' and w['a'][0] in defined and w['b'][0] in defined]
    if one:
        add('info', 'reciprocidade', f"{len(one)} fio(s) só estão escritos numa das pontas, embora as duas tenham folha: "
            + ', '.join(f"{':'.join(w['a'])}→{':'.join(w['b'])}" for w in one[:20]) + (' …' if len(one) > 20 else ''))
    skipped = [f'{sh} ({msg})' for k, sh, msg in problems if k == 'info']
    if skipped:
        add('info', 'folhas', 'Folhas não lidas: ' + '; '.join(skipped))
    return out


def topo_length(segments, a, b):
    adj = defaultdict(list)
    for x, y, l in segments:
        if l is None:
            continue
        adj[x].append((y, l))
        adj[y].append((x, l))
    if a not in adj or b not in adj:
        return None
    dist, q = {a: 0}, deque([a])
    while q:
        n = q.popleft()
        for m, l in adj[n]:
            if m not in dist:
                dist[m] = dist[n] + l
                q.append(m)
    return dist.get(b)


# ---------------------------------------------------------------------------- escrita do Excel
F = 'Arial'
HDR = dict(font=Font(name=F, bold=True, color='FFFFFF'), fill=PatternFill('solid', fgColor='1F2A44'))
BODY, PROP = Font(name=F), Font(name=F, italic=True, color='C55A11')
CALC = Font(name=F, color='7F7F7F')
TODO = PatternFill('solid', fgColor='FFF2CC')
N = 1000  # linhas cobertas pelas fórmulas


def key(ref, pin):
    # MATCH/COUNTIF ignoram maiúsculas e nos circulares A ≠ a: pinos com minúsculas levam "|lc"
    # shortcut: dois pinos que só diferem no padrão de maiúsculas ("aB" vs "Ab") colidiriam; não existem hoje
    return f'{ref}&":"&{pin}&IF(EXACT({pin},UPPER({pin})),"","|lc")'


def sheet(wb, title, cols, tab=None):
    ws = wb.create_sheet(title)
    for i, (name, width) in enumerate(cols, 1):
        c = ws.cell(1, i, name)
        c.font, c.fill = HDR['font'], HDR['fill']
        c.alignment = Alignment(vertical='center', wrap_text=True)
        ws.column_dimensions[c.column_letter].width = width
    ws.freeze_panes = 'A2'
    if tab:
        ws.sheet_properties.tabColor = tab
    return ws


def put(ws, r, c, v, todo=False, calc=False):
    cell = ws.cell(r, c, None if v == '' else v)
    cell.font = PROP if isinstance(v, P) else CALC if calc else BODY
    if todo and v in ('', None):
        cell.fill = TODO
    return cell


def finish(ws, check_col=None):
    ws.auto_filter.ref = ws.dimensions
    if check_col:
        ws.conditional_formatting.add(f'{check_col}2:{check_col}{N}', FormulaRule(
            formula=[f'LEN({check_col}2)>0'], font=Font(name=F, color='C00000', bold=True),
            fill=PatternFill('solid', fgColor='FCE4E4')))


def write(out, meta, pins, wires, cfg, issues, notes, studs, defined):
    code = meta['code']
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    readme = wb.create_sheet('README')
    segments = cfg.get('segments', [])

    # --- Wires
    ws = sheet(wb, 'Wires', [('wire_id', 12), ('harness', 9), ('net', 16), ('from_ref', 14), ('from_pin', 8),
                             ('to_ref', 14), ('to_pin', 8), ('colour', 8), ('csa_mm2', 8), ('awg', 6), ('hv', 5),
                             ('voltage', 8), ('wire_pn', 14), ('crimp_from_pn', 14), ('crimp_to_pn', 14),
                             ('twist_group', 9), ('shield', 7), ('length_mm', 9), ('description', 36),
                             ('origin', 20), ('issue', 10), ('check', 34), ('from_key', 14), ('to_key', 14),
                             ('wire_name', 14), ('notes', 30)], '2E75B6')
    for i, w in enumerate(wires, 2):
        L = cfg.get('wire_len', {}).get(frozenset((w['a'], w['b']))) or topo_length(segments, w['a'][0], w['b'][0])
        row = [f'W-{code}-{i - 1:03d}', code, w['net'], w['a'][0], w['a'][1], w['b'][0], w['b'][1],
               w['colour'], None, None, w['hv'], w['volt'], '', None, None, P(w['twist']) if w.get('twist') else '',
               '', L if L is not None else '', w['desc'], w['origin'], ', '.join(sorted(w['issues']))]
        for c, v in enumerate(row, 1):
            if v is not None:
                put(ws, i, c, v, todo=c in (8, 13, 18))
        if w['csa'] is not None:
            put(ws, i, 9, w['csa'])
        elif cfg.get('wire_default', {}).get('csa'):
            put(ws, i, 9, cfg['wire_default']['csa'])
        elif cfg.get('csa_default'):
            put(ws, i, 9, cfg['csa_default']).font = PROP  # proposta, ver README
        else:
            put(ws, i, 9, '', todo=True)
        if cfg.get('wire_default', {}).get('wire_pn'):
            put(ws, i, 13, cfg['wire_default']['wire_pn'])
        put(ws, i, 25, w['name'])
        put(ws, i, 26, w['note'])
        if w['colour_note']:
            ws.cell(i, 8).comment = Comment(w['colour_note'], 'migração')
        f = {10: f'=IF(I{i}="","",IFERROR(INDEX(Lists!$F$3:$F$17,MATCH(I{i},Lists!$E$3:$E$17,0)),"?"))',
             14: f'=IF(D{i}="","",IFERROR(INDEX(Connectors!$J$2:$J${N},MATCH(D{i},Connectors!$A$2:$A${N},0))&"",""))',
             15: f'=IF(F{i}="","",IFERROR(INDEX(Connectors!$J$2:$J${N},MATCH(F{i},Connectors!$A$2:$A${N},0))&"",""))',
             22: (f'=IF(D{i}="","",'
                  f'IF(COUNTIF(Connectors!$A$2:$A${N},D{i})=0,"from_ref não existe; ","")'
                  f'&IF(COUNTIF(Connectors!$A$2:$A${N},F{i})=0,"to_ref não existe; ","")'
                  f'&IF(AND(H{i}<>"",COUNTIF(Lists!$A$3:$A$14,H{i})=0),"cor inválida; ","")'
                  f'&IF(H{i}="","sem cor; ","")'
                  f'&IF((K{i}="Y")<>(H{i}="OG"),"HV e laranja não batem; ","")'
                  f'&IF(AND(K{i}="Y",OR(IFERROR(INDEX(Connectors!$F$2:$F${N},MATCH(D{i},Connectors!$A$2:$A${N},0)),"")<>"Y",'
                  f'IFERROR(INDEX(Connectors!$F$2:$F${N},MATCH(F{i},Connectors!$A$2:$A${N},0)),"")<>"Y")),"fio HV em ficha não HV rated; ","")'
                  f'&IF(AND(COUNTIF($W$2:$W${N},W{i})+COUNTIF($X$2:$X${N},W{i})>1,'
                  f'ISERROR(MATCH(IFERROR(INDEX(Connectors!$D$2:$D${N},MATCH(D{i},Connectors!$A$2:$A${N},0)),""),{{"SPLICE","STUD"}},0))),"pino from com >1 fio; ","")'
                  f'&IF(AND(COUNTIF($W$2:$W${N},X{i})+COUNTIF($X$2:$X${N},X{i})>1,'
                  f'ISERROR(MATCH(IFERROR(INDEX(Connectors!$D$2:$D${N},MATCH(F{i},Connectors!$A$2:$A${N},0)),""),{{"SPLICE","STUD"}},0))),"pino to com >1 fio; ","")'
                  f'&IF(I{i}="","sem secção; ",""))'),
             23: f'=IF(D{i}="","",{key(f"D{i}", f"E{i}")})', 24: f'=IF(F{i}="","",{key(f"F{i}", f"G{i}")})'}
        for c, v in f.items():
            put(ws, i, c, v, calc=True)
        if w['issues']:
            ws.cell(i, 1).fill = ws.cell(i, 21).fill = PatternFill('solid', fgColor='FBE2D5')
    finish(ws, 'V')
    for col, opts in (('H', '=Lists!$A$3:$A$14'), ('K', '"Y,N"'), ('Q', '"Y,N"')):
        dv = DataValidation(type='list', formula1=opts, allow_blank=True)
        dv.add(f'{col}2:{col}{N}')
        ws.add_data_validation(dv)

    # --- Connectors
    known, anchors = cfg.get('known', {}), cfg.get('anchors', {})
    refs = list(OrderedDict.fromkeys(r for r, _ in pins))
    wc = sheet(wb, 'Connectors', [('ref', 16), ('device', 10), ('harness', 9), ('type', 11), ('n_pins', 7),
                                  ('hv_rated', 8), ('housing_pn', 17), ('mate_pn', 14), ('manufacturer', 12),
                                  ('contact_pn', 13), ('seal_pn', 10), ('anchor_3d', 30), ('x_m', 7), ('y_m', 7),
                                  ('z_m', 7), ('in_topology', 9), ('n_wires', 8), ('check', 34), ('notes', 60),
                                  ('gender', 8)], '548235')
    for i, ref in enumerate(refs, 2):
        k = known.get(ref, {})
        typ = 'SPLICE' if ref.startswith('SPL') else 'STUD' if ref in studs else 'CONNECTOR'
        xyz = anchors.get(k.get('anchor'), ('', '', ''))
        note = k.get('notes', '') or ('' if ref in defined or typ == 'SPLICE' else 'Sem folha própria — criado a partir dos destinos.')
        row = [ref, ref.split()[0], code, typ, None, k.get('hv_rated', ''), k.get('housing_pn', ''),
               k.get('mate_pn', ''), k.get('manufacturer', ''), k.get('contact_pn', ''), '', k.get('anchor', ''),
               None, None, None, None, None, None, note]
        for c, v in enumerate(row, 1):
            if v is not None:
                put(wc, i, c, v, todo=typ != 'SPLICE' and c in (6, 7))
        # cada linha é UMA ficha (metade do conector): M ou F; a outra metade está noutro harness ou na placa
        put(wc, i, 20, k.get('gender', ''), todo=typ == 'CONNECTOR')
        for c, v in zip((13, 14, 15), xyz):  # centro da peça no GLB = proposta
            cell = put(wc, i, c, v, todo=typ != 'SPLICE')
            if v != '':
                cell.font = PROP
        put(wc, i, 5, f'=COUNTIF(Pins!$A$2:$A${N},A{i})', calc=True)
        put(wc, i, 16, f'=IF(COUNTIF(Segments!$B$2:$C${N},A{i})>0,"Y","N")', calc=True)
        put(wc, i, 17, f'=COUNTIF(Wires!$D$2:$D${N},A{i})+COUNTIF(Wires!$F$2:$F${N},A{i})', calc=True)
        put(wc, i, 18, f'=IF(A{i}="","",IF(D{i}="SPLICE","",IF(P{i}="N","fora da topologia; ","")'
                       f'&IF(G{i}="","sem P/N; ","")&IF(AND(G{i}<>"",COUNTIF(Parts!$A$2:$A${N},G{i})=0),"P/N não está em Parts; ","")'
                       f'&IF(M{i}="","sem xyz; ","")))', calc=True)
    finish(wc, 'R')
    for col, opts in (('D', '"CONNECTOR,STUD,SPLICE"'), ('F', '"Y,N"'), ('T', '"M,F"')):
        dv = DataValidation(type='list', formula1=opts, allow_blank=True)
        dv.add(f'{col}2:{col}{N}')
        wc.add_data_validation(dv)

    # --- Pins (gerado a partir de Wires)
    wp = sheet(wb, 'Pins', [('ref', 16), ('pin', 6), ('label', 16), ('dir', 7), ('voltage', 8), ('wire_id', 12),
                            ('net', 16), ('connects_to', 18), ('colour', 7), ('csa_mm2', 8), ('n_wires', 8),
                            ('check', 30), ('key', 14), ('_row', 6), ('_side', 6)], 'BF8F00')
    for i, ((ref, pin), info) in enumerate(pins.items(), 2):
        for c, v in enumerate([ref, pin, info['label'], info['dir'], info['volt']], 1):
            put(wp, i, c, v)
        if info['issue']:
            wp.cell(i, 3).comment = Comment(f"Label corrigido na migração ({info['issue']})", 'migração')
            wp.cell(i, 3).font = PROP
        f = {13: f'={key(f"A{i}", f"B{i}")}',
             14: f'=IFERROR(MATCH(M{i},Wires!$W$2:$W${N},0),IFERROR(MATCH(M{i},Wires!$X$2:$X${N},0),""))',
             15: f'=IF(N{i}="","",IF(ISNUMBER(MATCH(M{i},Wires!$W$2:$W${N},0)),"from","to"))',
             6: f'=IF(N{i}="","",INDEX(Wires!$A$2:$A${N},N{i}))',
             7: f'=IF(N{i}="","",INDEX(Wires!$C$2:$C${N},N{i}))',
             8: (f'=IF(N{i}="","",IF(O{i}="from",INDEX(Wires!$F$2:$F${N},N{i})&":"&INDEX(Wires!$G$2:$G${N},N{i}),'
                 f'INDEX(Wires!$D$2:$D${N},N{i})&":"&INDEX(Wires!$E$2:$E${N},N{i})))'),
             9: f'=IF(N{i}="","",INDEX(Wires!$H$2:$H${N},N{i})&"")',
             10: f'=IF(N{i}="","",IF(INDEX(Wires!$I$2:$I${N},N{i})="","",INDEX(Wires!$I$2:$I${N},N{i})))',
             11: f'=COUNTIF(Wires!$W$2:$W${N},M{i})+COUNTIF(Wires!$X$2:$X${N},M{i})',
             12: (f'=IF(AND(K{i}=0,C{i}<>"",D{i}<>"NC"),"pino com sinal mas sem fio; ","")'
                  f'&IF(AND(D{i}="NC",K{i}>0),"pino NC com fio; ","")'
                  f'&IF(AND(K{i}>1,ISERROR(MATCH(IFERROR(INDEX(Connectors!$D$2:$D${N},MATCH(A{i},Connectors!$A$2:$A${N},0)),""),{{"SPLICE","STUD"}},0))),"vários fios no mesmo pino; ","")')}
        for c, v in f.items():
            put(wp, i, c, v, calc=True)
    finish(wp, 'L')

    # --- Parts
    wq = sheet(wb, 'Parts', [('pn', 18), ('manufacturer', 12), ('type', 18), ('description', 55), ('awg_range', 11),
                             ('source', 26), ('confirmed', 9), ('datasheet', 30)], '7030A0')
    for i, row in enumerate(cfg.get('parts', []), 2):
        for c, v in enumerate(row, 1):
            put(wq, i, c, v)
        put(wq, i, 8, '', todo=True)
    finish(wq)

    # --- Segments / Nodes
    wsg = sheet(wb, 'Segments', [('seg_id', 8), ('from_node', 14), ('to_node', 14), ('length_mm', 10), ('harness', 9)], 'A5A5A5')
    for i, (a, b, l) in enumerate(segments, 2):
        for c, v in enumerate([f'S{i - 1:02d}', a, b, '' if l is None else l, code], 1):
            put(wsg, i, c, v, todo=c == 4)
    finish(wsg)
    wn = sheet(wb, 'Nodes', [('node', 8), ('x_m', 8), ('y_m', 8), ('z_m', 8), ('notes', 60)], 'A5A5A5')
    refs_ = {r for r, _ in pins}
    for i, n in enumerate(sorted({n for a, b, _ in segments for n in (a, b) if n not in refs_}), 2):
        note = ('ponto de ramificação do RapidHarness; xyz a medir no 3D (modo pick)' if re.fullmatch(r'BP\d+', n)
                else 'nome do desenho RapidHarness sem conector com este nome no Excel')
        for c, v in enumerate([n, '', '', '', note], 1):
            put(wn, i, c, v, todo=c in (2, 3, 4))
    finish(wn)

    # --- Issues
    wi = sheet(wb, 'Issues', [('id', 6), ('severity', 9), ('where', 22), ('problem', 70), ('option_a', 40),
                              ('option_b', 40), ('applied_here', 40), ('team_decision', 30), ('status', 9),
                              ('n_wires', 8)], 'C00000')
    for i, row in enumerate(issues, 2):
        for c, v in enumerate(row, 1):
            put(wi, i, c, v).alignment = Alignment(wrap_text=True, vertical='top')
        put(wi, i, 8, '', todo=True)
        put(wi, i, 9, 'aberto')
        put(wi, i, 10, f'=COUNTIF(Wires!$U$2:$U${N},"*"&A{i}&"*")', calc=True)
    finish(wi)
    dv = DataValidation(type='list', formula1='"aberto,decidido,fechado"', allow_blank=True)
    dv.add(f'I2:I{N}')
    wi.add_data_validation(dv)

    # --- Notes (folhas "Notas" do Excel original)
    wt = sheet(wb, 'Notes', [('sheet', 14), ('note', 100)], 'A5A5A5')
    for i, (sh, t) in enumerate(notes, 2):
        put(wt, i, 1, sh)
        put(wt, i, 2, t)

    # --- Lists
    wl = sheet(wb, 'Lists', [('colour', 8), ('nome', 11), ('uso proposto', 18), ('', 3), ('csa_mm2', 9), ('awg ≈', 7)], 'A5A5A5')
    put(wl, 2, 1, 'Cores IEC 60757')
    put(wl, 2, 5, 'Secção → AWG (equivalência aproximada)')
    for i, row in enumerate(COLOURS, 3):
        for c, v in enumerate(row, 1):
            put(wl, i, c, v)
    for i, (mm, awg) in enumerate(CSA_AWG, 3):
        put(wl, i, 5, mm)
        put(wl, i, 6, awg)

    # --- README
    readme.column_dimensions['A'].width = 24
    readme.column_dimensions['B'].width = 110
    lines = [
        (f"{CAR} — {meta['name']}: formato standard", None),
        ('Versão', f"v0.1 — gerado em {date.today():%Y-%m-%d} a partir de {', '.join(meta['sources'])} (tools/migrate.py)"),
        ('', ''),
        ('Regras', ''),
        ('1', 'Cada fio físico = 1 linha em Wires. Nunca se escreve o mesmo fio nas duas pontas.'),
        ('2', 'Pins é calculada a partir de Wires: diz, para cada pino, que fio lá está e para onde vai. '
              'Só as colunas ref/pin/label/dir/voltage se editam à mão.'),
        ('3', 'ref = nome do conector igual ao do RapidHarness (ex.: "AMS J11"). Pinos distinguem maiúsculas '
              '(C ≠ c nos circulares). Um conector com o mesmo nome noutro harness é a outra metade da ficha.'),
        ('4', 'Cores em código IEC 60757 (folha Lists). Laranja (OG) ⇔ hv = Y.'),
        ('5', 'Coluna check vazia = nada detetado. A vermelho = corrigir.'),
        ('6', 'Um pino com vários destinos leva um splice (ref "SPL …", type SPLICE). '
              'Vários olhais no mesmo perno = type STUD.'),
        ('', ''),
        ('Legenda', ''),
        ('preto', 'valor que veio do Excel original'),
        ('itálico laranja', 'valor PROPOSTO na migração — confirmar e passar a preto'),
        ('fundo amarelo', 'em falta — preencher'),
        ('cinzento', 'fórmula — não editar'),
        ('linha com issue', 'wire_id a salmão: a ligação foi corrigida ou tem dúvida — ver folha Issues'),
        ('', ''),
        ('Folhas', ''),
        ('Wires', 'lista de fios (a fonte da verdade)'),
        ('Connectors', 'uma linha por FICHA (metade do conector: gender M ou F), perno ou splice: P/N, crimp, HV rated, posição 3D. '
                       'Um conector (ex.: LVBOX J1) = ficha macho + ficha fêmea; cada harness tem a sua, ou a outra é o header da placa.'),
        ('Pins', 'pinout de cada ficha, gerado a partir de Wires'),
        ('Parts', 'biblioteca de P/N (fichas, contactos, cabos) com origem de cada um'),
        ('Segments / Nodes', 'topologia do RapidHarness: troços e comprimentos; xyz dos pontos de ramificação'),
        ('Issues', 'problemas encontrados no Excel original — a equipa decide e fecha'),
        ('Notes', 'notas da equipa que estavam nas folhas "Notas" do Excel original'),
        ('Lists', 'cores e secções válidas'),
    ]
    for r, (a, b) in enumerate(lines, 1):
        ca, cb = readme.cell(r, 1, a), readme.cell(r, 2, b)
        ca.font = Font(name=F, bold=True, size=14 if r == 1 else 10)
        cb.font = BODY
        cb.alignment = Alignment(wrap_text=True, vertical='top')
        style = {'itálico laranja': ('font', PROP), 'fundo amarelo': ('fill', TODO), 'cinzento': ('font', CALC),
                 'linha com issue': ('fill', PatternFill('solid', fgColor='FBE2D5'))}.get(a)
        if style:
            setattr(cb, *style)
    wb.calculation.fullCalcOnLoad = True
    wb.save(out)


# ---------------------------------------------------------------------------- um harness / todos
def with_proposals(cfg, meta, pins):
    """Junta às fichas as propostas da lista da equipa (catálogo); o que a config já define tem prioridade."""
    cat = config_parts.load()
    count = Counter(r for r, _ in pins)
    known = {}
    for ref in count:
        if ref.startswith('SPL'):
            continue
        prop = config_parts.propose(cat, meta['code'], ref, count[ref], meta['interface'], HUB, meta['hub_ifaces'])
        mine = cfg.get('known', {}).get(ref, {})
        if prop:
            notes = ' '.join(x for x in (mine.get('notes', ''), prop['notes']) if x)
            known[ref] = {**{k: P(v) for k, v in prop.items() if k != 'notes' and v}, **mine, 'notes': notes}
        elif mine:
            known[ref] = mine
    used = {v for k in known.values() for f, v in k.items() if f in ('housing_pn', 'contact_pn') and v}
    listed = {row[0] for row in cfg.get('parts', [])}
    parts = list(cfg.get('parts', [])) + config_parts.parts_rows(cat, used - listed)
    return dict(cfg, known=known, parts=parts)


def drawing_issues(segments, pins, issues):
    """Topologia vinda de um desenho: nomes que não existem no Excel e troços sem comprimento."""
    refs = {r for r, _ in pins}
    odd = sorted({n for a, b, _ in segments for n in (a, b) if n not in refs and not re.fullmatch(r'BP\d+', n)})
    nolen = [f'{a}–{b}' for a, b, l in segments if l is None]
    out = list(issues)
    if odd:
        out.append((f'I-{len(out) + 1:02d}', 'importante', 'desenho RapidHarness',
                    f"O desenho tem {len(odd)} nome(s) que não existem no Excel: {', '.join(odd)}.",
                    'Corrigir o Excel', 'Corrigir o desenho', 'Ficaram como nós do desenho (folha Nodes).'))
    if nolen:
        out.append((f'I-{len(out) + 1:02d}', 'menor', 'desenho RapidHarness',
                    f"{len(nolen)} troço(s) sem comprimento no desenho: {', '.join(nolen)}.", '', '', 'Comprimento por preencher'))
    return out


def migrate(path, meta, cfg):
    pins, edges, notes, problems, defined = read_workbook(path, cfg)
    cfg = with_proposals(cfg, meta, pins)
    studs = set(cfg.get('studs', set())) | {e['b'][0] for e in edges if e['b'][0] not in defined and ground_like(e['b'][0])}
    wires = build_wires(pins, edges, cfg, problems, studs)
    if cfg.get('strict'):
        self_check(edges, wires, cfg, studs)
        issues = cfg.get('issues', [])
    else:
        issues = auto_issues(pins, wires, problems, defined, studs)
    if cfg.get('segments') and not cfg.get('strict'):
        issues = drawing_issues(cfg['segments'], pins, issues)
    write(meta['out'], meta, pins, wires, cfg, issues, notes, studs, defined)
    return dict(wires=len(wires), connectors=len({r for r, _ in pins}), pins=len(pins), issues=len(issues)), list(defined)


def car_links(index, defined):
    """Acoplamentos HUB ↔ caixa (mesmo nome, prefixo de interface) e harnesses que parecem cópias uns dos outros."""
    car = [h for h in index['harnesses'] if h['category'] == 'car' and h['file']]
    mates = [dict(ref=ref, a=HUB, b=h['code']) for h in car if h['code'] != HUB
             for ref in defined[h['code']] if any(ref.startswith(p) for p in h['interface']) and ref in defined.get(HUB, [])]
    copies = []
    for i, x in enumerate(car):
        for y in car[i + 1:]:
            if HUB in (x['code'], y['code']):
                continue
            shared = sorted(set(defined[x['code']]) & set(defined[y['code']]))
            if shared and len(shared) >= 0.5 * min(len(defined[x['code']]), len(defined[y['code']])):
                copies.append(dict(a=x['code'], b=y['code'], refs=shared))
    # interface só de um lado: o HUB liga a ADBOX J1 mas o harness ADBOX não o tem (ou ao contrário)
    lonely = [dict(ref=ref, has=who, missing=other) for h in car if h['code'] != HUB and h['interface']
              for who, other, refs in ((HUB, h['code'], defined.get(HUB, [])), (h['code'], HUB, defined[h['code']]))
              for ref in refs if any(ref.startswith(p) for p in h['interface']) and ref not in defined[other if other != HUB else HUB]]
    return mates, copies, lonely


def main(src, dst):
    src, dst = Path(src), Path(dst)
    dst.mkdir(parents=True, exist_ok=True)
    index = dict(car=CAR, hub=HUB, generated=f'{date.today():%Y-%m-%d}', harnesses=[])
    defined = {}
    hub_ifaces = {p: code for code, _, cat, pre in HARNESSES.values() if cat == 'car' for p in pre}
    for folder in sorted(p for p in src.iterdir() if p.is_dir()):
        xlsx = sorted(f for f in folder.glob('*.xlsx') if not f.name.startswith('~$'))
        others = sorted(f.name for f in folder.iterdir() if f.is_file() and f.suffix.lower() != '.xlsx')
        code, name, category, iface = HARNESSES.get(folder.name, (re.sub(r'[^A-Z0-9]', '', folder.name.upper()), folder.name, 'car', []))
        if not xlsx:
            index['harnesses'].append(dict(code=code, name=name, folder=folder.name, category=category, file=None,
                                           interface=iface, sources=others))
            print(f'{code:13} sem Excel')
            continue
        for n, f in enumerate(xlsx):
            c, nm, cat, ifc = FILES.get(f.name, (code if len(xlsx) == 1 else f'{code}-{n + 1}', name, category, iface))
            meta = dict(code=c, name=nm, sources=[f.name], out=dst / f'{c}.xlsx', interface=ifc, hub_ifaces=hub_ifaces)
            cfg = CONFIGS[folder.name](P) if folder.name in CONFIGS else {}
            stats, defined[c] = migrate(f, meta, cfg)
            index['harnesses'].append(dict(code=c, name=nm, folder=folder.name, category=cat, file=f'{c}.xlsx',
                                           interface=ifc, sources=[f.name] + others, curated=bool(cfg), **stats))
            print(f"{c:13} {stats['wires']:4} fios {stats['connectors']:4} conectores {stats['pins']:4} pinos {stats['issues']:3} issues")
    index['mates'], index['copies'], index['lonely'] = car_links(index, defined)
    for m in index['mates']:
        print(f"acoplado: {m['ref']:18} {m['a']} ↔ {m['b']}")
    for l in index['lonely']:
        print(f"só num lado: {l['ref']:18} existe em {l['has']}, falta em {l['missing']}")
    for c in index['copies']:
        print(f"parecem cópias: {c['a']} e {c['b']} ({len(c['refs'])} conectores iguais)")
    (dst / 'index.json').write_text(json.dumps(index, ensure_ascii=False, indent=1))


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
