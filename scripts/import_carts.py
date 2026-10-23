"""Importa carrinhos da Mouser/Avancee (.xls) e DigiKey (.xlsx) para o catalogo de pecas.

Uso: .venv/bin/python scripts/import_carts.py carrinho1.xls carrinho2.xlsx ...
Os carrinhos mais recentes sao gravados por ultimo, por isso o preco mais novo de cada peca fica.
"""
import sys
from datetime import date
from pathlib import Path

import xlrd
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from backend import app  # noqa: E402
from backend.carts import parse_cart  # noqa: E402


def rows(path):
    if path.suffix.lower() == ".xlsx":
        return [list(row) for row in load_workbook(path, read_only=True, data_only=True).worksheets[0].iter_rows(values_only=True)]
    sheet = xlrd.open_workbook(path).sheet_by_index(0)
    return [sheet.row_values(i) for i in range(sheet.nrows)]


def main(paths):
    carts = [parse_cart(rows(path), path.name) for path in map(Path, paths)]
    carts.sort(key=lambda cart: cart[0] or date.min)
    parts = {part["mpn"].upper(): part for _, cart in carts for part in cart}
    app.save_parts(list(parts.values()))
    print(f"{len(parts)} peças importadas de {len(carts)} carrinhos")


if __name__ == "__main__":
    main(sys.argv[1:])
