"""Le as folhas de compras e encomendas (carrinhos Mouser/DigiKey, modelos BATCH COMPRAS/PRODUCAO, listas soltas).

Cada folha tem colunas diferentes: procura a linha de cabecalho pelos nomes conhecidos e le as linhas por baixo.
"""
import re
import unicodedata
from datetime import date, datetime

from carts import USD_TO_EUR, cart_date, money

# Nomes de cabecalho (sem acentos, minusculas) de cada campo; o primeiro que existir na folha ganha.
COLUMNS = {
    "pn": ("mfr. no", "no do fabricante", "manufacturer part number", "reference", "referencia", "ref", "valor/referencia",
           "product code", "n de idh", "codigo", "nome ficheiro", "nome do ficheiro", "specific example"),
    "supplier_pn": ("mouser no", "no da mouser", "mouser part number", "part number", "codigo carlis"),
    "description": ("description", "descricao", "producto - descripcion", "produto-descricao", "peca", "componente",
                    "name", "nome", "product name", "linha"),
    "manufacturer": ("manufacturer", "fabricante"),
    "qty": ("order qty.", "quant. do pedido", "quantity", "quantidade", "qt", "qtd:", "uds", "quantity needed", "meters",
            "quantidade a pedir"),
    "price": ("price (eur)", "preco (eur)", "unit price", "price per unit", "p.v.p. euro unitario", "p.v.p. / ud", "preco venda", "price", "preco"),
    "total": ("ext.: (eur)", "extended price usd", "price total", "p.vp.euro total", "total", "valor", "total expected"),
    "system": ("sistema", "system", "departament", "departamento afeto"),
    "notes": ("notas", "notes", "notas / aplicacao", "obs.", "material", "tipo", "type", "package", "dimensoes",
              "wire gauge", "colour", "capacidade", "capacidad"),
    "link": ("link",),
}
SKIP_SHEETS = re.compile(r"template", re.I)
# Tipo pelo nome do ficheiro/folha quando a folha nao e um modelo BATCH (que diz COMPRAS ou PRODUCAO no titulo).
KINDS = ((r"lista de precos", "Lista de preços"), (r"existencias", "Existências"),
         (r"necessidades|componentes|solicitud", "Pedido"), (r"produ", "Produção"))


def plain(value):
    text = unicodedata.normalize("NFKD", str(value).replace("€", "euro").replace("º", "o")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", text).strip().lower()


def cell_text(value):
    if value is None or str(value).startswith("#VALUE"):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return re.sub(r"\s+", " ", str(value)).strip()


def find_header(rows):
    """Indice da linha de cabecalho e {campo: coluna}; precisa de uma quantidade ou de P/N + descricao."""
    for index, row in enumerate(rows[:30]):
        names = [plain(cell) for cell in row]
        columns = {}
        for field, aliases in COLUMNS.items():
            found = next((names.index(alias) for alias in aliases if alias in names), None)
            if found is not None:
                columns[field] = found
        # "peca" so conta como descricao quando ha quantidade (evita a linha "Pessoa Responsavel da Peca").
        if "qty" in columns and columns.keys() & {"pn", "supplier_pn", "description"} or {"pn", "description"} <= columns.keys():
            if field_notes := [i for i, name in enumerate(names) if name in COLUMNS["notes"]]:
                columns["notes"] = field_notes
            return index, columns
    return None, {}


def sheet_meta(rows):
    """Fornecedor, data e tipo dos modelos BATCH (titulo + fornecedor nas 2 primeiras linhas, linha "Data")."""
    meta = {}
    for row in rows[:12]:
        cells = [cell for cell in row if cell_text(cell)]
        if not cells:
            continue
        first = plain(cells[0])
        title = next((cell_text(cell) for cell in cells if "batch" in plain(cell)), "")
        if title:
            meta["kind"] = "Produção" if "produ" in plain(title) else "Compra"
            continue
        if "kind" in meta and "supplier" not in meta and len(cells) == 1:
            meta["supplier"] = cell_text(cells[0])
        if first.startswith("pessoa responsavel") and len(cells) > 1 and plain(cells[1]) != "telemovel":
            meta["person"] = cell_text(cells[1])
        if first in ("data", "previsao de chegada do stock") and len(cells) > 1:
            meta["date"] = to_date(cells[1]) or meta.get("date")
        if first == "ano" or (len(cells) > 3 and plain(cells[2]) in ("ano", "data")):
            year = cells[3] if first != "ano" else cells[1]
            if re.fullmatch(r"20\d\d", cell_text(year)):
                meta["year"] = cell_text(year)
    return meta


def to_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    match = re.search(r"\b(\d{1,2})[-./](\d{1,2})[-./](\d{2,4})\b", str(value))
    if match:
        day, month, year = map(int, match.groups())
        try:
            return date(year + 2000 if year < 100 else year, month, day)
        except ValueError:
            return None
    return None


def number(value):
    if value in (None, ""):
        return None
    # Celulas numericas chegam como "51.27" (texto da celula); money() le "51,27 €" e tiraria o ponto.
    if re.fullmatch(r"-?\d+\.\d+", str(value)):
        return float(value)
    try:
        return money(value)
    except ValueError:
        match = re.search(r"\d+(?:[.,]\d+)?", str(value))
        return float(match.group().replace(",", ".")) if match else None


def file_supplier(name):
    known = ("Mouser", "DigiKey", "Wurth", "Avancee", "Carlis", "Coficab", "Loctite", "Henkel", "Hidromarinha", "Traco Power", "Mauser", "IMPL", "JLSantos")
    return next((supplier for supplier in known if supplier.lower() in name.lower()), "")


def parse_sheet(rows, file_name, sheet_name):
    """Linhas de compra de uma folha. Devolve [] se a folha nao tiver uma tabela reconhecivel."""
    if SKIP_SHEETS.search(sheet_name):
        return []
    header, columns = find_header(rows)
    if header is None:
        return []
    names = [plain(cell) for cell in rows[header]]
    usd = any("usd" in name for name in names)
    meta = sheet_meta(rows[:header])
    when = meta.get("date") or cart_date(rows[:header])
    kind = meta.get("kind") or next((kind for pattern, kind in KINDS if re.search(pattern, plain(f"{file_name} {sheet_name}"))), "Compra")
    batch = sheet_name.strip() if "batch" in sheet_name.lower() else re.sub(r"\.(xlsx?|csv)$", "", file_name.rsplit("/", 1)[-1], flags=re.I)
    supplier = meta.get("supplier") or file_supplier(f"{file_name} {sheet_name}")
    out = []
    for row in rows[header + 1:]:
        get = lambda field: cell_text(row[columns[field]]) if field in columns and columns[field] < len(row) else ""
        pn, description = get("pn"), get("description")
        if not (pn or description or get("supplier_pn")) or plain(description or pn) in ("total", "total*", "subtotal") or pn in ("NA", "-") and not description:
            continue
        price, total = number(get("price")), number(get("total"))
        if usd:
            price, total = (round(value * USD_TO_EUR, 4) if value is not None else None for value in (price, total))
        notes = " · ".join(text for i in columns.get("notes", []) if i < len(row) and (text := cell_text(row[i])) and text not in ("-", "NA"))
        out.append({
            "kind": kind, "supplier": supplier, "batch": batch,
            "date": when.isoformat() if when else meta.get("year", ""),
            "pn": "" if pn in ("NA", "-") else pn, "supplier_pn": get("supplier_pn"), "manufacturer": get("manufacturer"),
            "description": description, "qty": number(get("qty")), "price": price, "total": total,
            "system": get("system"), "person": meta.get("person", ""), "notes": notes, "link": get("link") if get("link").startswith("http") else "",
            "file": file_name, "sheet": sheet_name.strip(),
        })
    return out


def read_sheets(path):
    """[(nome da folha, linhas)] de um .xlsx, .xls ou .csv."""
    import csv
    import xlrd
    from openpyxl import load_workbook
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        return [(sheet.title, [list(row) for row in sheet.iter_rows(values_only=True)]) for sheet in workbook.worksheets]
    if suffix == ".xls":
        workbook = xlrd.open_workbook(path)
        return [(sheet.name, [sheet.row_values(i) for i in range(sheet.nrows)]) for sheet in workbook.sheets()]
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    return [("csv", list(csv.reader(text.splitlines(), delimiter=";" if text.count(";") > text.count(",") else ",")))]


def pdf_text(path):
    import subprocess
    try:
        result = subprocess.run(["pdftotext", "-l", "5", str(path), "-"], capture_output=True, text=True, timeout=20)
    except (FileNotFoundError, subprocess.TimeoutExpired):
        return ""
    return re.sub(r"\s+", " ", result.stdout).strip()


def load_folder(folder):
    """Todas as linhas das folhas da pasta, sem repetidos, e um registo por PDF (desenho ou batch impresso)."""
    rows, seen, problems = [], set(), []
    files = sorted(path for path in folder.rglob("*") if path.is_file())
    pdfs = {path.stem.lower(): path.relative_to(folder).as_posix() for path in files if path.suffix.lower() == ".pdf"}
    for path in files:
        name = path.relative_to(folder).as_posix()
        if path.suffix.lower() == ".pdf":
            text = pdf_text(path)
            rows.append({"kind": "Documento", "supplier": "", "batch": "", "date": "", "pn": path.stem, "supplier_pn": "",
                         "manufacturer": "", "description": text[:160], "qty": None, "price": None, "total": None, "system": "",
                         "person": "", "notes": "", "link": "", "file": name, "sheet": "", "doc": name, "text": text[:4000]})
            continue
        if path.suffix.lower() not in (".xlsx", ".xls", ".csv"):
            continue
        try:
            sheets = read_sheets(path)
        except Exception as exc:  # um ficheiro estragado nao pode esconder os outros
            problems.append(f"{name}: {exc}")
            continue
        for sheet_name, sheet_rows in sheets:
            for row in parse_sheet(sheet_rows, name, sheet_name):
                # O mesmo batch aparece em varios livros (copias de folhas, ou o carrinho Avancee guardado duas vezes): fica so a primeira.
                key = (row["kind"], row["date"], row["pn"].lower(), row["description"].lower(), row["qty"])
                if key in seen:
                    continue
                seen.add(key)
                row["doc"] = pdfs.get(row["pn"].lower(), "")
                rows.append(row)
    return rows, problems
